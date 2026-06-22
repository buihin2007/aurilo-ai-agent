
from openpyxl import load_workbook
from deep_translator import GoogleTranslator
import time
import re
from pathlib import Path
#map glossary to dict
wb_glossary = load_workbook("docs/glossary_finnish_english.xlsx")
ws_glossary = wb_glossary['Glossary FI-EN']
mapping = {}
for row in ws_glossary.iter_rows(min_row=2,max_row=ws_glossary.max_row, min_col=1,max_col=2):
    mapping[row[0].value]=row[1].value
#label translator function
translator = GoogleTranslator(source ="fi",target="en")
CODE_RE = re.compile(r"^(\d{4,5})\s+(.+)$")          # GL account prefix, e.g. "3654 Käyttöomaisuuden..."
def translate_label(input: str):
    if not isinstance(input, str) or not input.strip():
        return input
    text = input.strip()
    
    if text in mapping:
        return mapping[text]
    
    m = CODE_RE.match(text)
    if m:
        code, term = m.group(1), m.group(2)
        if term in mapping:
            return f"{code} {mapping[term]}"
    
    if "-" in text:
        head, tail = text.rsplit("-", 1)
        if tail in mapping:
            return f"{head}-{mapping[tail]}"
    
    output = translator.translate(text)
    time.sleep(0.3)
    if output:
        mapping[text] = output
    return output or input
#translate the excel files
def translate_file(path: Path):
    path = Path(path)
    wb = load_workbook(path, data_only =True)
    worksheets = wb.worksheets
    for worksheet in worksheets:
        for row in worksheet.iter_rows(min_row=1,max_row=worksheet.max_row, min_col=1,max_col=worksheet.max_column):
            for cell in row:
                cell_value = cell.value
                if isinstance(cell_value, str) and cell_value.strip():
                    translated = translate_label(cell_value)
                    if translated:
                        cell.value = translated
                else:
                    continue
    wb.save(path.parent / (path.stem+"_translated"+path.suffix))

